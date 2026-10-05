import datetime as dt
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(r"C:\xampp\htdocs\hrm\DUTY ROASTER - TEMPLATE.xlsx")
MYSQL = Path(r"C:\xampp\mysql\bin\mysql.exe")


def norm_name(value):
    value = re.sub(r"[^a-z0-9]+", "", str(value or "").lower())
    return value


def norm_dept(value):
    value = str(value or "").lower().replace("&", "and")
    return re.sub(r"[^a-z0-9]+", "", value)


def is_accounts_text(value):
    normalized = norm_dept(value)
    return "account" in normalized


def is_accounts_employee(employee):
    return is_accounts_text(employee.get("department", "")) or is_accounts_text(employee.get("position", ""))


def parse_date(value, year=2026):
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raw = str(value).strip()
    if raw in {"", "-"}:
        return None
    match = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", raw)
    if match:
        day, month, yy = match.groups()
        full_year = int(yy)
        if full_year < 100:
            full_year += 2000
        return dt.date(full_year, int(month), int(day))
    return None


def parse_time(value):
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, dt.time):
        return value.replace(microsecond=0)
    raw = str(value).strip()
    if raw in {"", "-"}:
        return None
    if re.match(r"^\d{1,2}:\d{2}(:\d{2})?$", raw):
        parts = [int(p) for p in raw.split(":")]
        while len(parts) < 3:
            parts.append(0)
        return dt.time(parts[0], parts[1], parts[2])
    return None


def sheet_date_range(title):
    dates = re.findall(r"(\d{1,2})[.](\d{1,2})[.](\d{4})", title)
    if len(dates) < 2:
        return None, None
    start = dt.date(int(dates[0][2]), int(dates[0][1]), int(dates[0][0]))
    end = dt.date(int(dates[1][2]), int(dates[1][1]), int(dates[1][0]))
    return start, end


def daterange(start, end):
    day = start
    while day <= end:
        yield day
        day += dt.timedelta(days=1)


def load_employees():
    sql = "SELECT id, employee_code, first_name, last_name, department, position FROM employees ORDER BY id"
    result = subprocess.run(
        [str(MYSQL), "-u", "root", "-P", "3307", "-h", "127.0.0.1", "--batch", "--raw", "hrmodule", "-e", sql],
        check=True,
        text=True,
        capture_output=True,
    )
    rows = []
    for line in result.stdout.splitlines()[1:]:
        cols = line.split("\t")
        if len(cols) < 6:
            continue
        rows.append({
            "id": int(cols[0]),
            "employee_code": cols[1],
            "name": f"{cols[2]} {cols[3]}",
            "department": cols[4],
            "position": cols[5],
        })
    return rows


def scalar_sql(sql):
    result = subprocess.run(
        [str(MYSQL), "-u", "root", "-P", "3307", "-h", "127.0.0.1", "--batch", "--raw", "--skip-column-names", "hrmodule", "-e", sql],
        check=True,
        text=True,
        capture_output=True,
    )
    return result.stdout.strip().splitlines()[0] if result.stdout.strip() else ""


def parse_roster():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    records = []
    for ws in wb.worksheets:
        start, end = sheet_date_range(ws.title)
        if not start or not end:
            continue
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8, values_only=True):
            sno, old_ec, name, dept, week_off, shift, start_time, end_time = row
            if not isinstance(sno, (int, float)) or not old_ec or not name or not shift:
                continue
            start_t = parse_time(start_time)
            end_t = parse_time(end_time)
            if not start_t or not end_t:
                continue
            records.append({
                "sheet": ws.title,
                "old_ec": str(old_ec).strip(),
                "name": str(name).strip(),
                "department": str(dept or "").strip(),
                "week_off": parse_date(week_off, start.year),
                "shift": str(shift).strip().upper(),
                "start": start_t.strftime("%H:%M:%S"),
                "end": end_t.strftime("%H:%M:%S"),
                "range_start": start,
                "range_end": end,
            })
    return records


def best_match(record, employees):
    rn = norm_name(record["name"])
    rd = norm_dept(record["department"])
    matches = [e for e in employees if norm_name(e["name"]) == rn]
    if len(matches) == 1:
        return matches[0]
    dept_matches = [e for e in matches if rd and rd in norm_dept(e["department"]) or norm_dept(e["department"]) in rd]
    if len(dept_matches) == 1:
        return dept_matches[0]
    return None


def sql_quote(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def assignment_rows(matched):
    rows = []
    for item in matched:
        for day in daterange(item["range_start"], item["range_end"]):
            is_off = item["week_off"] == day
            force_accounts_general = is_accounts_text(item["department"]) or is_accounts_employee(item["employee"])
            shift_name = item["shift"]
            start_time = item["start"]
            end_time = item["end"]
            if force_accounts_general:
                shift_name = "GS"
                start_time = "10:00:00"
                end_time = "18:00:00"
            rows.append({
                "employee_id": item["employee"]["id"],
                "duty_date": day.isoformat(),
                "shift_name": "OFF" if is_off else shift_name,
                "start_time": "00:00:00" if is_off else start_time,
                "end_time": "00:00:00" if is_off else end_time,
                "ward": item["department"] or item["employee"]["department"],
                "notes": f"Imported from DUTY ROASTER - TEMPLATE.xlsx; old EC {item['old_ec']}; sheet row employee {item['name']}; shift {item['shift']}",
            })
    return rows


def apply_assignments(matched):
    rows = assignment_rows(matched)
    if not rows:
        return {"deleted_scope": 0, "inserted": 0}

    creator_id = scalar_sql("SELECT id FROM users WHERE role='SuperAdmin' ORDER BY id LIMIT 1") or scalar_sql("SELECT id FROM users WHERE role='Admin' ORDER BY id LIMIT 1") or "1"
    employee_ids = sorted({row["employee_id"] for row in rows})
    dates = sorted({row["duty_date"] for row in rows})

    statements = [
        "START TRANSACTION;",
        "DELETE FROM duty_roster WHERE employee_id IN (" + ",".join(str(i) for i in employee_ids) + ") AND duty_date BETWEEN " + sql_quote(dates[0]) + " AND " + sql_quote(dates[-1]) + ";",
    ]
    for row in rows:
        statements.append(
            "INSERT INTO duty_roster(employee_id,duty_date,shift_name,start_time,end_time,ward,notes,created_by,created_at,updated_at) VALUES("
            + ",".join([
                str(row["employee_id"]),
                sql_quote(row["duty_date"]),
                sql_quote(row["shift_name"]),
                sql_quote(row["start_time"]),
                sql_quote(row["end_time"]),
                sql_quote(row["ward"]),
                sql_quote(row["notes"]),
                str(int(creator_id)),
                "NOW()",
                "NOW()",
            ])
            + ");"
        )
    statements.append("COMMIT;")

    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False, encoding="utf-8") as sql_file:
        sql_file.write("\n".join(statements))
        sql_path = sql_file.name
    try:
        subprocess.run(
            [str(MYSQL), "-u", "root", "-P", "3307", "-h", "127.0.0.1", "hrmodule", "--default-character-set=utf8mb4", "-e", f"source {sql_path}"],
            check=True,
            text=True,
            capture_output=True,
        )
    finally:
        Path(sql_path).unlink(missing_ok=True)

    inserted = int(scalar_sql(
        "SELECT COUNT(*) FROM duty_roster WHERE employee_id IN ("
        + ",".join(str(i) for i in employee_ids)
        + ") AND duty_date BETWEEN "
        + sql_quote(dates[0])
        + " AND "
        + sql_quote(dates[-1])
        + " AND notes LIKE 'Imported from DUTY ROASTER - TEMPLATE.xlsx%'"
    ) or "0")
    return {
        "matched_employees": len(employee_ids),
        "date_start": dates[0],
        "date_end": dates[-1],
        "inserted": inserted,
        "expected_inserted": len(rows),
    }


def main():
    employees = load_employees()
    records = parse_roster()
    matched = []
    unmatched = []
    for record in records:
        employee = best_match(record, employees)
        if employee:
            matched.append({**record, "employee": employee})
        else:
            unmatched.append(record)
    rows = assignment_rows(matched)
    total_rows = len(records)
    total_assignments = len(rows)
    off_assignments = sum(1 for row in rows if row["shift_name"] == "OFF")
    applied = apply_assignments(matched) if "--apply" in sys.argv else None
    print(json.dumps({
        "workbook": str(WORKBOOK),
        "mode": "apply" if "--apply" in sys.argv else "dry-run",
        "records": total_rows,
        "matched_records": len(matched),
        "unmatched_records": len(unmatched),
        "total_roster_days_to_write": total_assignments,
        "off_days_to_write": off_assignments,
        "applied": applied,
        "date_ranges": sorted({f"{r['range_start']} to {r['range_end']}" for r in records}),
        "unmatched": unmatched[:50],
        "sample_matched": [
            {
                "old_ec": r["old_ec"],
                "sheet_name": r["name"],
                "app_employee_code": r["employee"]["employee_code"],
                "app_name": r["employee"]["name"],
                "shift": r["shift"],
                "time": f"{r['start']} - {r['end']}",
                "week_off": str(r["week_off"] or "-"),
            }
            for r in matched[:20]
        ],
    }, indent=2, default=str))


if __name__ == "__main__":
    main()
